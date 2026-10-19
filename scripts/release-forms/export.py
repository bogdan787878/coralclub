"""Build the release-form workbook from data/*.jsonl — one tab per country."""
import json,os,re,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
countries=json.load(open('countries.json'))
LANG={'us':'English','ru':'Русский','de':'Deutsch','fr':'Français','es':'Español','it':'Italiano','pl':'Polski','cs':'Čeština','sk':'Slovenčina','hu':'Magyar','bg':'Български','ro':'Română','lv':'Latviešu','lt':'Lietuvių','et':'Eesti','tr':'Türkçe','ua':'Українська','il':'עברית'}
wb=Workbook(); summ=wb.active; summ.title='Сводка'
hdr_fill=PatternFill('solid',fgColor='004144'); hdr_font=Font(bold=True,color='FFFFFF')
def header(ws,row,vals):
    for i,v in enumerate(vals,1):
        c=ws.cell(row=row,column=i,value=v); c.fill=hdr_fill; c.font=hdr_font
used=set()
def sheet_name(n):
    s=re.sub(r'[\[\]:*?/\\]','',n)[:31] or 'Sheet'
    base,k=s,2
    while s.lower() in used: s=f"{base[:28]}-{k}"; k+=1
    used.add(s.lower()); return s
header(summ,1,['Страна','Домен','Язык','Продуктов','Без формы выпуска','Лист'])
r=2
for c in sorted(countries,key=lambda x:x['country']):
    path=f"data/{c['host']}.jsonl"
    if not os.path.exists(path): continue
    recs=[json.loads(l) for l in open(path,encoding='utf8')]
    recs=[x for x in recs if not x.get('error')]
    if not recs: continue
    recs.sort(key=lambda x:(x['name'] or x['slug']).lower())
    name=sheet_name(c['country']); ws=wb.create_sheet(name)
    header(ws,1,['Продукт','Название поля','Значения (форма выпуска)','Ссылка','Slug'])
    row=2; nofm=0
    for x in recs:
        blocks=x['blocks'] or [['—',[]]]
        if not x['blocks']: nofm+=1
        for head,vals in blocks:
            ws.cell(row=row,column=1,value=x['name'] or x['slug'])
            ws.cell(row=row,column=2,value=head)
            ws.cell(row=row,column=3,value=' | '.join(vals) if vals else '—')
            ws.cell(row=row,column=4,value=x['url']).hyperlink=x['url']
            ws.cell(row=row,column=5,value=x['slug'])
            row+=1
    for col,w in zip('ABCDE',[48,22,60,60,36]): ws.column_dimensions[col].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=f"A1:E{row-1}"
    summ.cell(row=r,column=1,value=c['country']); summ.cell(row=r,column=2,value=c['host'])
    summ.cell(row=r,column=3,value=LANG.get(c['lang'],c['lang'])); summ.cell(row=r,column=4,value=len(recs))
    summ.cell(row=r,column=5,value=nofm)
    cl=summ.cell(row=r,column=6,value=name); cl.hyperlink=f"#'{name}'!A1"; cl.font=Font(color='0563C1',underline='single')
    r+=1
for col,w in zip('ABCDEF',[28,26,14,12,20,28]): summ.column_dimensions[col].width=w
summ.freeze_panes='A2'
out=sys.argv[1] if len(sys.argv)>1 else 'release-forms.xlsx'
wb.save(out); print('saved',out,len(wb.sheetnames)-1,'country tabs')
